"""Excel workbook export / import.

The workbook is generated from data/catalog.json after every change made on the
website. Derived columns (unit cost, recipe cost, reorder status, usage...) are
live Excel formulas, so the sheet keeps working if you edit numbers in Excel.

Edits made in Excel can be pulled back into the website with the import
(website "Excel" tab, or `python -m app.excel import <file>`). Only the
blue input columns of the Ingredients, Menu Items and Recipes sheets are read.
"""

import os
import re
import sys
from datetime import date, time

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .store import ROOT, UNITS, Store, week_start

EXPORT_PATH = os.path.join(ROOT, "exports", "inventory.xlsx")

FONT = "Arial"
INPUT = Font(name=FONT, color="0000FF")
CALC = Font(name=FONT, color="000000")
LINK = Font(name=FONT, color="008000")
HEADER = Font(name=FONT, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5D50")
INPUT_HEADER_FILL = PatternFill("solid", fgColor="3E7CB1")
TITLE = Font(name=FONT, bold=True, size=14)
THIN = Border(bottom=Side(style="thin", color="D0D7DE"))
RED_FILL = PatternFill("solid", fgColor="F8D7DA")
GREEN_FILL = PatternFill("solid", fgColor="D4EDDA")
MATRIX_FILL = PatternFill("solid", fgColor="E2F0D9")

MONEY = '$#,##0.00;($#,##0.00);"-"'
MONEY4 = '$#,##0.0000;($#,##0.0000);"-"'
QTY = '#,##0.##;-#,##0.##;"-"'
PCT = '0.0%;-0.0%;"-"'

# (header, key, kind, width, number_format)
#   kind: "in" = user input read back on import, "f" = formula, "link" = formula to other sheet
ING_COLS = [
    ("Ingredient ID", "id", "in", 12, None),
    ("Name", "name", "in", 22, None),
    ("Category", "category", "in", 13, None),
    ("Recipe Unit", "unit", "in", 10, None),
    ("Supplier", "supplier", "in", 20, None),
    ("Supplier SKU", "supplier_sku", "in", 13, None),
    ("Package Description", "package_desc", "in", 20, None),
    ("Qty per Package (units)", "package_qty", "in", 12, QTY),
    ("Package Cost ($)", "package_cost", "in", 11, MONEY),
    ("Shipping per Package ($)", "shipping_cost", "in", 11, MONEY),
    ("Landed Unit Cost ($)", None, "f", 12, MONEY4),
    ("Lead Time (days)", "lead_time_days", "in", 9, "0"),
    ("Shelf Life (days)", "shelf_life_days", "in", 9, "0"),
    ("Min Order (pkgs)", "min_order_packages", "in", 9, "0"),
    ("On Hand (units)", "on_hand", "in", 11, QTY),
    ("Reorder Point (units)", "reorder_point", "in", 11, QTY),
    ("Par Level (units)", "par_level", "in", 11, QTY),
    ("Used in # Items", None, "f", 9, "0"),
    ("Units Used (sales log)", None, "f", 11, QTY),
    ("Status", None, "f", 10, None),
    ("Suggested Order (pkgs)", None, "f", 11, "0"),
    ("Order Cost ($)", None, "f", 11, MONEY),
    ("Stock Value ($)", None, "f", 11, MONEY),
    ("Notes", "notes", "in", 30, None),
]
ITEM_COLS = [
    ("Item ID", "id", "in", 10, None),
    ("Name", "name", "in", 24, None),
    ("Category", "category", "in", 14, None),
    ("Size", "size", "in", 9, None),
    ("Price ($)", "price", "in", 10, MONEY),
    ("Description", "description", "in", 44, None),
    ("Image File", "image", "in", 18, None),
    ("Active", "active", "in", 8, None),
    ("# Ingredients", None, "f", 10, "0"),
    ("Ingredient Cost ($)", None, "f", 12, MONEY),
    ("Gross Margin ($)", None, "f", 12, MONEY),
    ("Cost %", None, "f", 9, PCT),
    ("Units Sold", None, "f", 9, "#,##0"),
    ("Revenue ($)", None, "f", 12, MONEY),
]
RECIPE_COLS = [
    ("Item ID", "item_id", "in", 10, None),
    ("Item Name", None, "link", 24, None),
    ("Ingredient ID", "ingredient_id", "in", 12, None),
    ("Ingredient Name", None, "link", 22, None),
    ("Qty per Item", "qty", "in", 10, QTY),
    ("Unit", None, "link", 8, None),
    ("Unit Cost ($)", None, "link", 11, MONEY4),
    ("Line Cost ($)", None, "f", 11, MONEY),
    ("Item Units Sold", None, "link", 10, "#,##0"),
    ("Ingredient Used", None, "f", 11, QTY),
]


def col(cols, header):
    """Column letter for a header in one of the column specs above."""
    return get_column_letter([c[0] for c in cols].index(header) + 1)


def _header_row(ws, cols, row=1):
    for idx, (header, _key, kind, width, _fmt) in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=idx, value=header)
        cell.font = HEADER
        cell.fill = INPUT_HEADER_FILL if kind == "in" else HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[row].height = 42
    ws.freeze_panes = ws.cell(row=row + 1, column=3)


def _style_cell(cell, kind, fmt):
    cell.font = {"in": INPUT, "f": CALC, "link": LINK}[kind]
    cell.border = THIN
    if fmt:
        cell.number_format = fmt


# ---------------------------------------------------------------- export
def export_workbook(store: Store, path=EXPORT_PATH):
    data = store.data
    ings, items, sales = data["ingredients"], data["items"], data["sales"]
    recipe_rows = [(it, line) for it in items for line in it.get("recipe", [])]
    wb = Workbook()

    # Sheet sizes: give formulas room to grow so rows added in Excel are counted.
    n_ing = max(len(ings) + 50, 200)
    n_items = max(len(items) + 50, 200)
    n_rec = max(len(recipe_rows) + 200, 1000)
    n_sales = max(len(sales) + 500, 2000)
    ING = f"Ingredients!$A$2:$A${n_ing}"
    ITM = f"'Menu Items'!$A$2:$A${n_items}"

    def ing_lookup(header, key_cell):
        c = col(ING_COLS, header)
        return f'=IFERROR(INDEX(Ingredients!${c}$2:${c}${n_ing},MATCH({key_cell},{ING},0)),"")'

    def item_lookup(header, key_cell):
        c = col(ITEM_COLS, header)
        return f"=IFERROR(INDEX('Menu Items'!${c}$2:${c}${n_items},MATCH({key_cell},{ITM},0)),\"\")"

    # --- Guide -------------------------------------------------------------
    ws = wb.active
    ws.title = "Guide"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100
    ws["A1"] = "Menu & Parts Inventory Workbook"
    ws["A1"].font = TITLE
    guide = [
        ("How it works", "Each menu item is built from ingredients (its recipe / bill of materials). "
                         "Selling an item consumes those ingredients, which drives the reorder list."),
        ("Ingredients", "Master list of every part you buy: purchasing info, stock, reorder rules."),
        ("Menu Items", "What you sell. Cost, margin and units sold are calculated from the recipes."),
        ("Recipes", "One row per ingredient per menu item. To add an ingredient to an item, add a row "
                    "with the Item ID, Ingredient ID and quantity."),
        ("Where Used", "Matrix of ingredients vs. menu items - quickly see which items need a part."),
        ("Procurement", "Reorder list grouped by supplier. Rows flagged REORDER should be purchased."),
        ("Purchase Orders", "Orders created from the website (history)."),
        ("Sales Log", "Sales recorded on the website; feeds 'Units Used' and 'Units Sold'."),
        ("Staff", "Employees with role, hourly rate and maximum hours per week."),
        ("Schedule", "Every shift: date, employee, start/end, hours and labor cost. OPEN = not yet assigned."),
        ("Weekly Hours", "Hours per employee per day for one week. Change the Monday date in B1 to "
                         "look at another week; over-max rows turn red."),
        ("", ""),
        ("Colour legend", "Blue header / blue text = input you can edit.  Dark green header = calculated, "
                          "do not overwrite.  Green text = pulled from another sheet."),
        ("Units", "All stock quantities are in the ingredient's Recipe Unit (e.g. oz). "
                  "'Qty per Package' says how many recipe units are in one purchased package."),
        ("Landed Unit Cost", "(Package Cost + Shipping per Package) / Qty per Package."),
        ("Reorder rule", "When On Hand <= Reorder Point, order enough packages to refill to Par Level "
                         "(at least the supplier's Min Order). Tip: Reorder Point ~= daily usage x lead "
                         "time + safety stock."),
        ("Syncing", "The website rewrites this file after every change. Edits made here can be pulled "
                    "into the website from its Excel tab (Import). Only blue input columns are read."),
    ]
    for r, (a, b) in enumerate(guide, start=3):
        ws.cell(row=r, column=1, value=a).font = Font(name=FONT, bold=True)
        c = ws.cell(row=r, column=2, value=b)
        c.font = Font(name=FONT)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    # --- Ingredients -------------------------------------------------------
    ws = wb.create_sheet("Ingredients")
    _header_row(ws, ING_COLS)
    R = {h: col(RECIPE_COLS, h) for h, *_ in RECIPE_COLS}
    for r, ing in enumerate(ings, start=2):
        L = {h: f"{col(ING_COLS, h)}{r}" for h, *_ in ING_COLS}
        formulas = {
            "Landed Unit Cost ($)": f"=IFERROR(({L['Package Cost ($)']}+{L['Shipping per Package ($)']})"
                                    f"/{L['Qty per Package (units)']},0)",
            "Used in # Items": f"=COUNTIF(Recipes!${R['Ingredient ID']}$2:${R['Ingredient ID']}${n_rec},"
                               f"{L['Ingredient ID']})",
            "Units Used (sales log)": f"=SUMIFS(Recipes!${R['Ingredient Used']}$2:${R['Ingredient Used']}${n_rec},"
                                      f"Recipes!${R['Ingredient ID']}$2:${R['Ingredient ID']}${n_rec},"
                                      f"{L['Ingredient ID']})",
            "Status": f'=IF({L["On Hand (units)"]}<={L["Reorder Point (units)"]},"REORDER","OK")',
            "Suggested Order (pkgs)": f"=IF({L['On Hand (units)']}<={L['Reorder Point (units)']},"
                                      f"MAX({L['Min Order (pkgs)']},ROUNDUP(MAX({L['Par Level (units)']}"
                                      f"-{L['On Hand (units)']},0)/{L['Qty per Package (units)']},0)),0)",
            "Order Cost ($)": f"={L['Suggested Order (pkgs)']}*({L['Package Cost ($)']}"
                              f"+{L['Shipping per Package ($)']})",
            "Stock Value ($)": f"=MAX({L['On Hand (units)']},0)*{L['Landed Unit Cost ($)']}",
        }
        for c, (header, key, kind, _w, fmt) in enumerate(ING_COLS, start=1):
            cell = ws.cell(row=r, column=c, value=ing.get(key) if key else formulas[header])
            _style_cell(cell, kind, fmt)
    last = len(ings) + 1
    status = col(ING_COLS, "Status")
    ws.conditional_formatting.add(
        f"A2:{get_column_letter(len(ING_COLS))}{n_ing}",
        FormulaRule(formula=[f'${status}2="REORDER"'], fill=RED_FILL))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(ING_COLS))}{last}"
    dv = DataValidation(type="list", formula1='"' + ",".join(UNITS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col(ING_COLS, 'Recipe Unit')}2:{col(ING_COLS, 'Recipe Unit')}{n_ing}")
    ws[f"{col(ING_COLS, 'Reorder Point (units)')}1"].comment = Comment(
        "Reorder when On Hand falls to this level. Rule of thumb: daily usage x lead time + safety stock.",
        "Inventory")

    # --- Menu Items --------------------------------------------------------
    ws = wb.create_sheet("Menu Items")
    _header_row(ws, ITEM_COLS)
    for r, it in enumerate(items, start=2):
        L = {h: f"{col(ITEM_COLS, h)}{r}" for h, *_ in ITEM_COLS}
        formulas = {
            "# Ingredients": f"=COUNTIF(Recipes!${R['Item ID']}$2:${R['Item ID']}${n_rec},{L['Item ID']})",
            "Ingredient Cost ($)": f"=SUMIFS(Recipes!${R['Line Cost ($)']}$2:${R['Line Cost ($)']}${n_rec},"
                                   f"Recipes!${R['Item ID']}$2:${R['Item ID']}${n_rec},{L['Item ID']})",
            "Gross Margin ($)": f"={L['Price ($)']}-{L['Ingredient Cost ($)']}",
            "Cost %": f"=IFERROR({L['Ingredient Cost ($)']}/{L['Price ($)']},0)",
            "Units Sold": f"=SUMIFS('Sales Log'!$C$2:$C${n_sales},'Sales Log'!$B$2:$B${n_sales},{L['Item ID']})",
            "Revenue ($)": f"={L['Units Sold']}*{L['Price ($)']}",
        }
        for c, (header, key, kind, _w, fmt) in enumerate(ITEM_COLS, start=1):
            value = formulas[header] if not key else ("Yes" if key == "active" and it.get(key) else
                                                       "No" if key == "active" else it.get(key))
            cell = ws.cell(row=r, column=c, value=value)
            _style_cell(cell, kind, fmt)
            if header == "Description":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(ITEM_COLS))}{len(items) + 1}"
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col(ITEM_COLS, 'Active')}2:{col(ITEM_COLS, 'Active')}{n_items}")
    pct = col(ITEM_COLS, "Cost %")
    ws.conditional_formatting.add(f"{pct}2:{pct}{n_items}",
                                  CellIsRule(operator="greaterThan", formula=["0.35"], fill=RED_FILL))

    # --- Recipes (bill of materials) ----------------------------------------
    ws = wb.create_sheet("Recipes")
    _header_row(ws, RECIPE_COLS)
    for r, (it, line) in enumerate(recipe_rows, start=2):
        L = {h: f"{col(RECIPE_COLS, h)}{r}" for h, *_ in RECIPE_COLS}
        formulas = {
            "Item Name": item_lookup("Name", L["Item ID"]),
            "Ingredient Name": ing_lookup("Name", L["Ingredient ID"]),
            "Unit": ing_lookup("Recipe Unit", L["Ingredient ID"]),
            "Unit Cost ($)": ing_lookup("Landed Unit Cost ($)", L["Ingredient ID"]).replace(',"")', ",0)"),
            "Line Cost ($)": f"=IFERROR({L['Qty per Item']}*{L['Unit Cost ($)']},0)",
            "Item Units Sold": item_lookup("Units Sold", L["Item ID"]).replace(',"")', ",0)"),
            "Ingredient Used": f"=IFERROR({L['Qty per Item']}*{L['Item Units Sold']},0)",
        }
        values = {"item_id": it["id"], "ingredient_id": line["ingredient_id"], "qty": line["qty"]}
        for c, (header, key, kind, _w, fmt) in enumerate(RECIPE_COLS, start=1):
            cell = ws.cell(row=r, column=c, value=values[key] if key else formulas[header])
            _style_cell(cell, kind, fmt)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(RECIPE_COLS))}{len(recipe_rows) + 1}"
    for header, rng in (("Item ID", ITM), ("Ingredient ID", ING)):
        dv = DataValidation(type="list", formula1=rng, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col(RECIPE_COLS, header)}2:{col(RECIPE_COLS, header)}{n_rec}")

    # --- Where Used matrix ---------------------------------------------------
    ws = wb.create_sheet("Where Used")
    ws["A1"] = "Ingredient"
    ws["B1"] = "Ingredient ID"
    ws["C1"] = "Unit"
    ws["A2"] = "Quantity of each ingredient per menu item (blank = not used)"
    ws["A2"].font = Font(name=FONT, italic=True, color="666666")
    for c in ("A1", "B1", "C1"):
        ws[c].font, ws[c].fill = HEADER, HEADER_FILL
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 8
    rid, ring, rqty = R["Item ID"], R["Ingredient ID"], R["Qty per Item"]
    for c, it in enumerate(items, start=4):
        letter = get_column_letter(c)
        ws.column_dimensions[letter].width = 11
        h = ws.cell(row=1, column=c, value=f"={item_lookup('Name', f'{letter}$2')[1:]}")
        h.font, h.fill = HEADER, HEADER_FILL
        h.alignment = Alignment(text_rotation=90, horizontal="center")
        ws.cell(row=2, column=c, value=it["id"]).font = Font(name=FONT, bold=True)
    ws.row_dimensions[1].height = 130
    for r, ing in enumerate(ings, start=3):
        ws.cell(row=r, column=1, value=f"={ing_lookup('Name', f'$B{r}')[1:]}").font = LINK
        ws.cell(row=r, column=2, value=ing["id"]).font = CALC
        ws.cell(row=r, column=3, value=f"={ing_lookup('Recipe Unit', f'$B{r}')[1:]}").font = LINK
        for c in range(4, len(items) + 4):
            letter = get_column_letter(c)
            cell = ws.cell(row=r, column=c,
                           value=f"=SUMIFS(Recipes!${rqty}$2:${rqty}${n_rec},Recipes!${rid}$2:${rid}${n_rec},"
                                 f"{letter}$2,Recipes!${ring}$2:${ring}${n_rec},$B{r})")
            cell.font, cell.number_format = CALC, '#,##0.##;-#,##0.##;""'
            cell.alignment = Alignment(horizontal="center")
    if items and ings:
        ws.conditional_formatting.add(
            f"D3:{get_column_letter(len(items) + 3)}{len(ings) + 2}",
            CellIsRule(operator="greaterThan", formula=["0"], fill=MATRIX_FILL))
    ws.freeze_panes = "D3"

    # --- Procurement -----------------------------------------------------------
    ws = wb.create_sheet("Procurement")
    proc_cols = [
        ("Supplier", "Supplier", 20, None), ("Ingredient", "Name", 22, None),
        ("Supplier SKU", "Supplier SKU", 13, None), ("Package", "Package Description", 20, None),
        ("On Hand", "On Hand (units)", 10, QTY), ("Reorder Point", "Reorder Point (units)", 10, QTY),
        ("Status", "Status", 10, None), ("Order (pkgs)", "Suggested Order (pkgs)", 10, "0"),
        ("Lead Time (days)", "Lead Time (days)", 10, "0"), ("Order Cost ($)", "Order Cost ($)", 12, MONEY),
    ]
    for c, (header, _src, width, _fmt) in enumerate(proc_cols, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font, cell.fill = HEADER, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.cell(row=1, column=len(proc_cols) + 1, value="Expected Arrival").font = HEADER
    ws.cell(row=1, column=len(proc_cols) + 1).fill = HEADER_FILL
    ws.column_dimensions[get_column_letter(len(proc_cols) + 1)].width = 13
    ws.row_dimensions[1].height = 32
    order = sorted(range(len(ings)), key=lambda i: (ings[i]["supplier"].lower(), ings[i]["name"].lower()))
    for r, idx in enumerate(order, start=2):
        src_row = idx + 2
        for c, (_h, src, _w, fmt) in enumerate(proc_cols, start=1):
            cell = ws.cell(row=r, column=c, value=f"=Ingredients!{col(ING_COLS, src)}{src_row}")
            _style_cell(cell, "link", fmt)
        lead = f"I{r}"
        cell = ws.cell(row=r, column=len(proc_cols) + 1, value=f'=IF(G{r}="REORDER",TODAY()+{lead},"")')
        _style_cell(cell, "f", "yyyy-mm-dd")
    tot = len(order) + 3
    ws.cell(row=tot, column=1, value="Total to order").font = Font(name=FONT, bold=True)
    c = ws.cell(row=tot, column=10, value=f"=SUM(J2:J{tot - 2})")
    c.font, c.number_format = Font(name=FONT, bold=True), MONEY
    ws.conditional_formatting.add(f"A2:K{max(tot - 2, 2)}",
                                  FormulaRule(formula=['$G2="REORDER"'], fill=RED_FILL))
    ws.conditional_formatting.add(f"G2:G{max(tot - 2, 2)}",
                                  CellIsRule(operator="equal", formula=['"OK"'], fill=GREEN_FILL))
    ws.auto_filter.ref = f"A1:K{max(tot - 2, 2)}"
    ws.freeze_panes = "C2"

    # --- Purchase Orders -------------------------------------------------------
    ws = wb.create_sheet("Purchase Orders")
    po_cols = [("PO #", 9), ("Supplier", 20), ("Created", 11), ("Expected", 11), ("Status", 10),
               ("Ingredient ID", 12), ("Ingredient", 22), ("Packages", 9), ("Package Cost ($)", 11),
               ("Shipping ($)", 11), ("Line Total ($)", 12)]
    for c, (header, width) in enumerate(po_cols, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font, cell.fill = HEADER, HEADER_FILL
        ws.column_dimensions[get_column_letter(c)].width = width
    r = 2
    for po in data["purchase_orders"]:
        for line in po["lines"]:
            ing = store.ingredient(line["ingredient_id"]) or {}
            vals = [po["id"], po["supplier"], po["created"], po["expected"], po["status"],
                    line["ingredient_id"], ing.get("name", ""), line["packages"],
                    line["package_cost"], line["shipping_cost"], f"=H{r}*(I{r}+J{r})"]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=c, value=v)
                _style_cell(cell, "f" if c == 11 else "in", MONEY if c >= 9 else None)
            r += 1
    ws.freeze_panes = "A2"

    # --- Sales Log -------------------------------------------------------------
    ws = wb.create_sheet("Sales Log")
    for c, (header, width) in enumerate([("Date / Time", 20), ("Item ID", 10), ("Qty", 7),
                                         ("Item Name", 24), ("Revenue ($)", 12)], start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font, cell.fill = HEADER, HEADER_FILL
        ws.column_dimensions[get_column_letter(c)].width = width
    for r, s in enumerate(sales, start=2):
        _style_cell(ws.cell(row=r, column=1, value=s["date"].replace("T", " ")), "in", None)
        _style_cell(ws.cell(row=r, column=2, value=s["item_id"]), "in", None)
        _style_cell(ws.cell(row=r, column=3, value=s["qty"]), "in", "0")
        _style_cell(ws.cell(row=r, column=4, value=item_lookup("Name", f"B{r}")), "link", None)
        price = col(ITEM_COLS, "Price ($)")
        _style_cell(ws.cell(row=r, column=5,
                            value=f"=IFERROR(C{r}*INDEX('Menu Items'!${price}$2:${price}${n_items},"
                                  f"MATCH(B{r},{ITM},0)),0)"), "f", MONEY)
    ws.freeze_panes = "A2"

    _schedule_sheets(wb, store)

    for sheet in wb.worksheets:
        sheet.sheet_view.zoomScale = 100
    wb.calculation.fullCalcOnLoad = True

    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, path)
    return path


STAFF_COLS = [
    ("Employee ID", "id", "in", 12, None),
    ("Name", "name", "in", 20, None),
    ("Role", "role", "in", 12, None),
    ("Hourly Rate ($)", "hourly_rate", "in", 10, MONEY),
    ("Max Hours / Week", "max_hours", "in", 10, "0.#"),
    ("Phone", "phone", "in", 15, None),
    ("Email", "email", "in", 24, None),
    ("Active", "active", "in", 8, None),
]
SCHED_COLS = [
    ("Shift ID", None, "in", 10, None),
    ("Date", None, "in", 11, "yyyy-mm-dd"),
    ("Day", None, "f", 6, None),
    ("Employee ID", None, "in", 12, None),
    ("Employee", None, "link", 20, None),
    ("Role", None, "in", 11, None),
    ("Start", None, "in", 8, "hh:mm"),
    ("End", None, "in", 8, "hh:mm"),
    ("Hours", None, "f", 8, "0.0#"),
    ("Hourly Rate ($)", None, "link", 10, MONEY),
    ("Labor Cost ($)", None, "f", 11, MONEY),
    ("Week Of", None, "f", 11, "yyyy-mm-dd"),
    ("Notes", None, "in", 24, None),
]


def _schedule_sheets(wb, store):
    """Staff, Schedule (one row per shift) and a Weekly Hours grid driven by formulas."""
    emps = store.data["employees"]
    shifts = sorted(store.data["shifts"], key=lambda x: (x["date"], x["start"]))
    n_emp = max(len(emps) + 50, 100)
    n_sh = max(len(shifts) + 500, 2000)

    ws = wb.create_sheet("Staff")
    _header_row(ws, STAFF_COLS)
    for r, e in enumerate(emps, start=2):
        for c, (_h, key, kind, _w, fmt) in enumerate(STAFF_COLS, start=1):
            v = ("Yes" if e[key] else "No") if key == "active" else e.get(key)
            _style_cell(ws.cell(row=r, column=c, value=v), kind, fmt)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(STAFF_COLS))}{len(emps) + 1}"

    S = {h: col(SCHED_COLS, h) for h, *_ in SCHED_COLS}
    rate = col(STAFF_COLS, "Hourly Rate ($)")
    ws = wb.create_sheet("Schedule")
    _header_row(ws, SCHED_COLS)
    for r, sh in enumerate(shifts, start=2):
        L = {h: f"{S[h]}{r}" for h in S}
        hh, mm = map(int, sh["start"].split(":"))
        eh, em = map(int, sh["end"].split(":"))
        values = {
            "Shift ID": sh["id"],
            "Date": date.fromisoformat(sh["date"]),
            "Day": f'=TEXT({L["Date"]},"ddd")',
            "Employee ID": sh["employee_id"] or "OPEN",
            "Employee": f'=IF({L["Employee ID"]}="OPEN","OPEN SHIFT",IFERROR(INDEX(Staff!$B$2:$B${n_emp},'
                        f'MATCH({L["Employee ID"]},Staff!$A$2:$A${n_emp},0)),""))',
            "Role": sh["role"],
            "Start": time(hh, mm),
            "End": time(eh, em),
            "Hours": f"=MOD({L['End']}-{L['Start']},1)*24",
            "Hourly Rate ($)": f"=IFERROR(INDEX(Staff!${rate}$2:${rate}${n_emp},"
                               f"MATCH({L['Employee ID']},Staff!$A$2:$A${n_emp},0)),0)",
            "Labor Cost ($)": f"={L['Hours']}*{L['Hourly Rate ($)']}",
            "Week Of": f"={L['Date']}-WEEKDAY({L['Date']},3)",
            "Notes": sh.get("notes", ""),
        }
        for c, (header, _key, kind, _w, fmt) in enumerate(SCHED_COLS, start=1):
            _style_cell(ws.cell(row=r, column=c, value=values[header]), kind, fmt)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SCHED_COLS))}{len(shifts) + 1}"
    ws.conditional_formatting.add(f"A2:{get_column_letter(len(SCHED_COLS))}{n_sh}",
                                  FormulaRule(formula=[f'${S["Employee ID"]}2="OPEN"'],
                                              fill=PatternFill("solid", fgColor="FDF2DC")))

    # --- Weekly Hours grid ----------------------------------------------------
    ws = wb.create_sheet("Weekly Hours")
    ws["A1"] = "Week of"
    ws["A1"].font = Font(name=FONT, bold=True)
    ws["B1"] = week_start(date.today())
    ws["B1"].font, ws["B1"].number_format = INPUT, "yyyy-mm-dd"
    ws["B1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["C1"] = "Change this Monday date to view another week."
    ws["C1"].font = Font(name=FONT, italic=True, color="666666")
    heads = ["Employee ID", "Name", "Role"] + [None] * 7 + ["Total Hours", "Max Hours", "Over Max?", "Labor Cost ($)"]
    widths = [12, 20, 11] + [11] * 7 + [10, 9, 9, 12]
    for c, (h, w) in enumerate(zip(heads, widths), start=1):
        cell = ws.cell(row=3, column=c, value=h if h else f"=$B$1+{c - 4}")
        cell.font, cell.fill = HEADER, HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if not h:
            cell.number_format = "ddd m/d"
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[3].height = 32
    rng = lambda h: f"Schedule!${S[h]}$2:${S[h]}${n_sh}"
    roster = [e for e in emps if e["active"]] + [e for e in emps if not e["active"]]
    rows = [(e["id"], e["name"], e["role"], e["max_hours"]) for e in roster] + [("OPEN", "Open shifts", "", None)]
    for r, (eid, name, role, mx) in enumerate(rows, start=4):
        for c, v in enumerate([eid, name, role], start=1):
            _style_cell(ws.cell(row=r, column=c, value=v), "f" if c == 1 else "link", None)
        for c in range(4, 11):
            d = f"{get_column_letter(c)}$3"
            cell = ws.cell(row=r, column=c, value=f'=SUMIFS({rng("Hours")},{rng("Employee ID")},$A{r},{rng("Date")},{d})')
            _style_cell(cell, "f", '0.0#;-0.0#;""')
            cell.alignment = Alignment(horizontal="center")
        _style_cell(ws.cell(row=r, column=11, value=f"=SUM(D{r}:J{r})"), "f", "0.0#")
        _style_cell(ws.cell(row=r, column=12, value=mx), "link", "0.#")
        _style_cell(ws.cell(row=r, column=13, value=f'=IF(AND(L{r}<>"",K{r}>L{r}),"OVER","")'), "f", None)
        _style_cell(ws.cell(row=r, column=14,
                            value=f'=SUMIFS({rng("Labor Cost ($)")},{rng("Employee ID")},$A{r},'
                                  f'{rng("Week Of")},$B$1)'), "f", MONEY)
    last = len(rows) + 3
    tot = last + 1
    ws.cell(row=tot, column=2, value="Total hours").font = Font(name=FONT, bold=True)
    ws.cell(row=tot + 1, column=2, value="Labor cost ($)").font = Font(name=FONT, bold=True)
    for c in range(4, 12):
        L = get_column_letter(c)
        _style_cell(ws.cell(row=tot, column=c, value=f"=SUM({L}4:{L}{last})"), "f", "0.0#")
        ws.cell(row=tot, column=c).font = Font(name=FONT, bold=True)
        if c <= 10:
            v = f'=SUMIFS({rng("Labor Cost ($)")},{rng("Date")},{L}$3)'
            _style_cell(ws.cell(row=tot + 1, column=c, value=v), "f", '$#,##0')
    _style_cell(ws.cell(row=tot + 1, column=14, value=f"=SUM(N4:N{last})"), "f", MONEY)
    ws.cell(row=tot + 1, column=14).font = Font(name=FONT, bold=True)
    ws.conditional_formatting.add(f"A4:N{last}", FormulaRule(formula=['$M4="OVER"'], fill=RED_FILL))
    ws.conditional_formatting.add(f"D4:J{last}", CellIsRule(operator="greaterThan", formula=["0"], fill=MATRIX_FILL))
    ws.freeze_panes = "D4"


def export_listener(store):
    """Store listener: keep the workbook in sync after every save."""
    try:
        export_workbook(store)
    except PermissionError:
        return "Excel file is open in another program - close it so it can be updated."
    return None


# ---------------------------------------------------------------- import
def _rows(ws, cols):
    headers = [c.value for c in ws[1]]
    index = {}
    for header, key, kind, *_ in cols:
        if kind == "in":
            if header not in headers:
                raise ValueError(f"Sheet '{ws.title}' is missing column '{header}'")
            index[key] = headers.index(header)
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = {k: row[i] if i < len(row) else None for k, i in index.items()}
        if any(v not in (None, "") for v in rec.values()):
            yield rec


def read_workbook(path):
    """Parse the input columns of a workbook into (ingredients, items) payloads."""
    wb = load_workbook(path)  # formulas are ignored: we only read input columns
    for name in ("Ingredients", "Menu Items", "Recipes"):
        if name not in wb.sheetnames:
            raise ValueError(f"Workbook has no '{name}' sheet")
    ingredients = list(_rows(wb["Ingredients"], ING_COLS))
    items = list(_rows(wb["Menu Items"], ITEM_COLS))
    by_name = {str(i["name"]).strip().lower(): i for i in ingredients if i.get("name")}
    for ing in ingredients:
        ing["id"] = str(ing["id"]).strip() if ing.get("id") else None
    nums = [int(m.group(1)) for i in ingredients if i["id"] and (m := re.match(r"ING-(\d+)$", i["id"]))]
    nxt = max(nums, default=0)
    for ing in ingredients:
        if not ing["id"]:  # new row typed in Excel without an ID
            nxt += 1
            ing["id"] = f"ING-{nxt:03d}"
    for it in items:
        it["id"] = str(it["id"]).strip() if it.get("id") else None
        it["recipe"] = []
    item_map = {it["id"]: it for it in items if it["id"]}
    ing_ids = {ing["id"] for ing in ingredients}
    problems = []
    for n, line in enumerate(_rows(wb["Recipes"], RECIPE_COLS), start=2):
        item = item_map.get(str(line["item_id"]).strip() if line["item_id"] else None)
        ing_id = str(line["ingredient_id"]).strip() if line["ingredient_id"] else ""
        if ing_id not in ing_ids and ing_id.lower() in by_name:  # allow typing a name
            ing_id = by_name[ing_id.lower()]["id"]
        if not item:
            problems.append(f"Recipes row {n}: unknown Item ID '{line['item_id']}'")
        elif ing_id not in ing_ids:
            problems.append(f"Recipes row {n}: unknown Ingredient ID '{line['ingredient_id']}'")
        else:
            item["recipe"].append({"ingredient_id": ing_id, "qty": line["qty"]})
    if problems:
        raise ValueError("; ".join(problems[:10]))
    return ingredients, items


def import_workbook(store: Store, path):
    ingredients, items = read_workbook(path)
    store.replace_catalog(ingredients, items)
    return {"ingredients": len(ingredients), "items": len(items)}


if __name__ == "__main__":
    s = Store()
    if len(sys.argv) >= 3 and sys.argv[1] == "import":
        print(import_workbook(s, sys.argv[2]))
        s.save()
    else:
        print("Wrote", export_workbook(s, sys.argv[1] if len(sys.argv) > 1 else EXPORT_PATH))
